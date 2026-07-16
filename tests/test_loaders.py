"""Unit tests for src/loaders (Phase 2)."""

from __future__ import annotations

from pathlib import Path

import pytest

from src.loaders.document_loader import DocumentLoader
from src.loaders.docx_loader import DocxLoader
from src.loaders.markdown_loader import MarkdownLoader
from src.loaders.pdf_loader import PDFLoader
from src.loaders.txt_loader import TxtLoader


class TestTxtLoader:
    def test_loads_plain_text(self, tmp_path: Path) -> None:
        file_path = tmp_path / "note.txt"
        file_path.write_text("Hello, RAG system.", encoding="utf-8")

        document = TxtLoader().load(file_path)

        assert document is not None
        assert document["filename"] == "note.txt"
        assert document["text"] == "Hello, RAG system."
        assert document["metadata"]["extension"] == ".txt"

    def test_skips_empty_file(self, tmp_path: Path) -> None:
        file_path = tmp_path / "empty.txt"
        file_path.write_text("   \n\n", encoding="utf-8")

        assert TxtLoader().load(file_path) is None

    def test_handles_latin1_encoding_gracefully(self, tmp_path: Path) -> None:
        file_path = tmp_path / "latin1.txt"
        file_path.write_bytes("café résumé".encode("latin-1"))

        document = TxtLoader().load(file_path)

        assert document is not None
        assert "caf" in document["text"]


class TestMarkdownLoader:
    def test_loads_markdown_preserving_syntax(self, tmp_path: Path) -> None:
        file_path = tmp_path / "doc.md"
        file_path.write_text("# Title\n\nSome **bold** content.", encoding="utf-8")

        document = MarkdownLoader().load(file_path)

        assert document is not None
        assert "# Title" in document["text"]
        assert document["metadata"]["extension"] == ".md"


class TestDocxLoader:
    def test_loads_paragraphs_and_tables(self, tmp_path: Path) -> None:
        import docx

        file_path = tmp_path / "handbook.docx"
        doc = docx.Document()
        doc.add_paragraph("Employee Handbook")
        doc.add_paragraph("All employees get 20 PTO days per year.")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Role"
        table.rows[0].cells[1].text = "Engineer"
        doc.save(str(file_path))

        document = DocxLoader().load(file_path)

        assert document is not None
        assert "Employee Handbook" in document["text"]
        assert "20 PTO days" in document["text"]
        assert "Engineer" in document["text"]


class TestPDFLoader:
    def test_loads_text_and_counts_pages(self, tmp_path: Path) -> None:
        from reportlab.pdfgen import canvas

        file_path = tmp_path / "manual.pdf"
        pdf = canvas.Canvas(str(file_path))
        pdf.drawString(72, 750, "Warranty covers 12 months from purchase.")
        pdf.showPage()
        pdf.drawString(72, 750, "Contact support for claims.")
        pdf.showPage()
        pdf.save()

        document = PDFLoader().load(file_path)

        assert document is not None
        assert "Warranty" in document["text"]
        assert document["metadata"]["pages"] == 2
        assert document["metadata"]["page_texts"] is not None
        assert len(document["metadata"]["page_texts"]) == 2


class TestDocumentLoader:
    def test_loads_all_supported_types_and_skips_unsupported(self, tmp_path: Path) -> None:
        (tmp_path / "a.txt").write_text("text content", encoding="utf-8")
        (tmp_path / "b.md").write_text("# markdown content", encoding="utf-8")
        (tmp_path / "c.xyz").write_text("unsupported", encoding="utf-8")

        documents = DocumentLoader().load_directory(tmp_path)

        filenames = {d["filename"] for d in documents}
        assert filenames == {"a.txt", "b.md"}

    def test_recurses_into_subdirectories(self, tmp_path: Path) -> None:
        nested = tmp_path / "nested" / "deep"
        nested.mkdir(parents=True)
        (nested / "deep.txt").write_text("deep content", encoding="utf-8")

        documents = DocumentLoader().load_directory(tmp_path)

        assert any(d["filename"] == "deep.txt" for d in documents)

    def test_raises_for_missing_directory(self, tmp_path: Path) -> None:
        with pytest.raises(FileNotFoundError):
            DocumentLoader().load_directory(tmp_path / "does-not-exist")

    def test_empty_directory_returns_empty_list(self, tmp_path: Path) -> None:
        assert DocumentLoader().load_directory(tmp_path) == []
